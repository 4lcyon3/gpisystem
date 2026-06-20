import io
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import WorksheetProperties, PageSetupProperties


class PlantillaService:
    """Servicio para generar plantillas Excel profesionales con validaciones"""
    
    # Estilos corporativos
    HEADER_FILL = PatternFill(start_color='1e3a8a', end_color='1e3a8a', fill_type='solid')
    HEADER_FONT = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    REQUIRED_FILL = PatternFill(start_color='fef3c7', end_color='fef3c7', fill_type='solid')
    BORDER = Border(
        left=Side(style='thin', color='d1d5db'),
        right=Side(style='thin', color='d1d5db'),
        top=Side(style='thin', color='d1d5db'),
        bottom=Side(style='thin', color='d1d5db')
    )
    
    @staticmethod
    def _crear_hoja_instrucciones(wb: Workbook, titulo: str, campos: dict):
        """Crea una hoja de instrucciones con descripción de cada campo"""
        ws = wb.create_sheet("Instrucciones", 0)
        
        # Título
        ws['A1'] = f"📋 Instrucciones: {titulo}"
        ws['A1'].font = Font(size=14, bold=True, color='1e3a8a')
        ws.merge_cells('A1:D1')
        
        # Introducción
        ws['A3'] = "ℹ️ INSTRUCCIONES GENERALES"
        ws['A3'].font = Font(size=12, bold=True)
        ws['A4'] = "1. Complete la hoja 'Datos' con la información requerida"
        ws['A5'] = "2. Las columnas en AMARILLO son obligatorias"
        ws['A6'] = "3. Las columnas con dropdown (▼) tienen valores predefinidos"
        ws['A7'] = "4. No modifique los encabezados de las columnas"
        ws['A8'] = "5. Guarde el archivo como .xlsx (Excel)"
        
        # Descripción de campos
        ws['A10'] = "📝 DESCRIPCIÓN DE CAMPOS"
        ws['A10'].font = Font(size=12, bold=True)
        
        ws['A11'] = "Campo"
        ws['B11'] = "Obligatorio"
        ws['C11'] = "Tipo"
        ws['D11'] = "Descripción"
        for col in ['A11', 'B11', 'C11', 'D11']:
            ws[col].font = Font(bold=True, color='FFFFFF')
            ws[col].fill = PlantillaService.HEADER_FILL
        
        row = 12
        for campo, info in campos.items():
            ws[f'A{row}'] = campo
            ws[f'B{row}'] = "Sí" if info.get('required') else "No"
            ws[f'C{row}'] = info.get('type', 'Texto')
            ws[f'D{row}'] = info.get('description', '')
            row += 1
        
        # Ajustar anchos
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 12
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 60
    
    @staticmethod
    def _aplicar_estilo_headers(ws, num_columnas: int):
        """Aplica estilos corporativos a los headers"""
        for col in range(1, num_columnas + 1):
            cell = ws.cell(row=1, column=col)
            cell.fill = PlantillaService.HEADER_FILL
            cell.font = PlantillaService.HEADER_FONT
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = PlantillaService.BORDER
    
    @staticmethod
    def generar_centros_costo() -> io.BytesIO:
        """Genera plantilla Excel para Centros de Costo"""
        wb = Workbook()
        
        # Hoja de Datos
        ws = wb.active
        ws.title = "Datos" # type: ignore
        
        # Headers
        headers = [
            "ruc_entidad", "codigo", "nombre", 
            "nivel_jerarquico", "centro_padre", "activo"
        ]
        ws.append(headers) # type: ignore
        PlantillaService._aplicar_estilo_headers(ws, len(headers))
        
        # Fila de ejemplo
        ws.append([ # type: ignore
            "20131380014", "001", "Dirección General de Salud",
            "Nivel 1", "", "SI"
        ])
        ws.append([ # type: ignore
            "20131380014", "001.01", "Oficina de Planificación",
            "Nivel 2", "001", "SI"
        ])
        
        # Marcar campos obligatorios (primeras 3 columnas)
        for row in range(2, 4):
            for col in range(1, 4):
                ws.cell(row=row, column=col).fill = PlantillaService.REQUIRED_FILL # type: ignore
        
        # Dropdown para nivel_jerarquico (columna 4)
        dv_nivel = DataValidation(
            type="list",
            formula1='"Nivel 1,Nivel 2,Nivel 3,Nivel 4"',
            allow_blank=True
        )
        dv_nivel.add(f'D2:D1000')
        ws.add_data_validation(dv_nivel) # type: ignore
        
        # Dropdown para activo (columna 6)
        dv_activo = DataValidation(
            type="list",
            formula1='"SI,NO"',
            allow_blank=True
        )
        dv_activo.add(f'F2:F1000')
        ws.add_data_validation(dv_activo) # type: ignore
        
        # Validación de longitud para RUC (11 dígitos)
        dv_ruc = DataValidation(
            type="textLength",
            operator="equal",
            formula1="11"
        )
        dv_ruc.error = "El RUC debe tener exactamente 11 dígitos"
        dv_ruc.errorTitle = "RUC Inválido"
        dv_ruc.add(f'A2:A1000')
        ws.add_data_validation(dv_ruc) # type: ignore
        
        # Ajustar anchos
        ws.column_dimensions['A'].width = 15 # type: ignore
        ws.column_dimensions['B'].width = 12 # type: ignore
        ws.column_dimensions['C'].width = 35 # type: ignore
        ws.column_dimensions['D'].width = 18 # type: ignore
        ws.column_dimensions['E'].width = 15 # type: ignore
        ws.column_dimensions['F'].width = 10 # type: ignore
        
        # Hoja de Instrucciones
        campos = {
            "ruc_entidad": {
                "required": True,
                "type": "Número (11 dígitos)",
                "description": "RUC de la entidad a la que pertenece el centro de costo"
            },
            "codigo": {
                "required": True,
                "type": "Texto",
                "description": "Código único del centro de costo (ej: 001, 001.01)"
            },
            "nombre": {
                "required": True,
                "type": "Texto",
                "description": "Nombre descriptivo del centro de costo"
            },
            "nivel_jerarquico": {
                "required": False,
                "type": "Dropdown",
                "description": "Nivel en la estructura orgánica (Nivel 1 = Dirección, Nivel 2 = Oficina, etc.)"
            },
            "centro_padre": {
                "required": False,
                "type": "Texto",
                "description": "Código del centro de costo padre (para estructura jerárquica)"
            },
            "activo": {
                "required": False,
                "type": "Dropdown (SI/NO)",
                "description": "Indica si el centro de costo está activo (por defecto: SI)"
            },
        }
        PlantillaService._crear_hoja_instrucciones(wb, "Centros de Costo", campos)
        
        # Guardar
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output
    
    @staticmethod
    def generar_metas_presupuestales() -> io.BytesIO:
        """Genera plantilla Excel para Metas Presupuestales"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Datos" # type: ignore
        
        headers = [
            "ruc_entidad", "codigo", "nombre", "anio_fiscal",
            "tipo_meta", "unidad_medida", "activo"
        ]
        ws.append(headers) # type: ignore
        PlantillaService._aplicar_estilo_headers(ws, len(headers))
        
        # Ejemplos
        ws.append([ # type: ignore
            "20131380014", "00001", "Atención de salud a la población",
            2026, "Producto", "Atenciones", "SI"
        ])
        ws.append([ # type: ignore
            "20131380014", "00002", "Construcción Hospital Regional",
            2026, "Proyecto", "Metros cuadrados", "SI"
        ])
        
        # Marcar obligatorios
        for row in range(2, 4):
            for col in range(1, 5):
                ws.cell(row=row, column=col).fill = PlantillaService.REQUIRED_FILL # type: ignore
        
        # Dropdowns
        dv_tipo = DataValidation(
            type="list",
            formula1='"Producto,Proyecto,Actividad,Inversión"',
            allow_blank=True
        )
        dv_tipo.add(f'E2:E1000')
        ws.add_data_validation(dv_tipo) # type: ignore
        
        dv_activo = DataValidation(type="list", formula1='"SI,NO"', allow_blank=True)
        dv_activo.add(f'G2:G1000')
        ws.add_data_validation(dv_activo) # type: ignore
        
        # Validación de año (2000-2100)
        dv_anio = DataValidation(type="whole", operator="between", formula1="2000", formula2="2100")
        dv_anio.error = "El año debe estar entre 2000 y 2100"
        dv_anio.add(f'D2:D1000')
        ws.add_data_validation(dv_anio) # type: ignore
        
        # Anchuras
        for i, width in enumerate([15, 12, 40, 12, 15, 20, 10], 1):
            ws.column_dimensions[get_column_letter(i)].width = width # type: ignore
        
        campos = {
            "ruc_entidad": {"required": True, "type": "Número (11 dígitos)", "description": "RUC de la entidad"},
            "codigo": {"required": True, "type": "Texto", "description": "Código único de la meta (ej: 00001)"},
            "nombre": {"required": True, "type": "Texto", "description": "Nombre descriptivo de la meta presupuestal"},
            "anio_fiscal": {"required": True, "type": "Número (2000-2100)", "description": "Año fiscal de la meta"},
            "tipo_meta": {"required": False, "type": "Dropdown", "description": "Tipo de meta: Producto, Proyecto, Actividad o Inversión"},
            "unidad_medida": {"required": False, "type": "Texto", "description": "Unidad de medida (ej: Atenciones, M², Informes)"},
            "activo": {"required": False, "type": "Dropdown (SI/NO)", "description": "Estado de la meta (por defecto: SI)"},
        }
        PlantillaService._crear_hoja_instrucciones(wb, "Metas Presupuestales", campos)
        
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output
    
    @staticmethod
    def generar_fuentes_financiamiento() -> io.BytesIO:
        """Genera plantilla Excel para Fuentes de Financiamiento"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Datos" # type: ignore
        
        headers = ["codigo", "nombre", "tipo_rubro", "descripcion_detallada", "activo"]
        ws.append(headers) # type: ignore
        PlantillaService._aplicar_estilo_headers(ws, len(headers))
        
        ws.append(["1", "Recursos Ordinarios", "RO", "Fondos del tesoro público", "SI"]) # type: ignore
        ws.append(["2", "Recursos Directamente Recaudados", "RDR", "Ingresos propios de la entidad", "SI"]) # type: ignore
        
        for row in range(2, 4):
            for col in range(1, 4):
                ws.cell(row=row, column=col).fill = PlantillaService.REQUIRED_FILL # type: ignore
        
        dv_tipo = DataValidation(
            type="list",
            formula1='"RO,RDR,ROOC,DT,RD,RE,IM,AC"',
            allow_blank=True
        )
        dv_tipo.add(f'C2:C1000')
        ws.add_data_validation(dv_tipo) # type: ignore
        
        dv_activo = DataValidation(type="list", formula1='"SI,NO"', allow_blank=True)
        dv_activo.add(f'E2:E1000')
        ws.add_data_validation(dv_activo) # type: ignore
        
        for i, width in enumerate([12, 40, 15, 50, 10], 1):
            ws.column_dimensions[get_column_letter(i)].width = width # type: ignore
        
        campos = {
            "codigo": {"required": True, "type": "Texto", "description": "Código único de la fuente (ej: 1, 2)"},
            "nombre": {"required": True, "type": "Texto", "description": "Nombre oficial de la fuente de financiamiento"},
            "tipo_rubro": {"required": True, "type": "Dropdown", "description": "RO=Recursos Ordinarios, RDR=Recursos Directamente Recaudados, etc."},
            "descripcion_detallada": {"required": False, "type": "Texto", "description": "Descripción extendida de la fuente"},
            "activo": {"required": False, "type": "Dropdown (SI/NO)", "description": "Estado de la fuente (por defecto: SI)"},
        }
        PlantillaService._crear_hoja_instrucciones(wb, "Fuentes de Financiamiento", campos)
        
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output
    
    @staticmethod
    def generar_clasificadores() -> io.BytesIO:
        """Genera plantilla Excel para Clasificadores de Gasto"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Datos" # type: ignore
        
        headers = [
            "codigo", "descripcion", "nivel",
            "clasificador_padre", "activo"
        ]
        ws.append(headers) # type: ignore
        PlantillaService._aplicar_estilo_headers(ws, len(headers))
        
        ws.append(["2.3", "Bienes y Servicios", "Genérica", "", "SI"]) # type: ignore
        ws.append(["2.3.2", "Bienes", "Subgenérica", "2.3", "SI"]) # type: ignore
        ws.append(["2.3.2.1", "Equipos médicos", "Específica", "2.3.2", "SI"]) # type: ignore
        
        for row in range(2, 5):
            for col in range(1, 3):
                ws.cell(row=row, column=col).fill = PlantillaService.REQUIRED_FILL # type: ignore
        
        dv_nivel = DataValidation(
            type="list",
            formula1='"Genérica,Subgenérica,Específica"',
            allow_blank=True
        )
        dv_nivel.add(f'C2:C1000')
        ws.add_data_validation(dv_nivel) # type: ignore
        
        dv_activo = DataValidation(type="list", formula1='"SI,NO"', allow_blank=True)
        dv_activo.add(f'E2:E1000')
        ws.add_data_validation(dv_activo) # type: ignore
        
        for i, width in enumerate([15, 45, 15, 20, 10], 1):
            ws.column_dimensions[get_column_letter(i)].width = width # type: ignore
        
        campos = {
            "codigo": {"required": True, "type": "Texto", "description": "Código del clasificador (ej: 2.3, 2.3.2.1)"},
            "descripcion": {"required": True, "type": "Texto", "description": "Descripción del clasificador de gasto"},
            "nivel": {"required": False, "type": "Dropdown", "description": "Nivel jerárquico: Genérica, Subgenérica o Específica"},
            "clasificador_padre": {"required": False, "type": "Texto", "description": "Código del clasificador padre (para estructura jerárquica)"},
            "activo": {"required": False, "type": "Dropdown (SI/NO)", "description": "Estado del clasificador (por defecto: SI)"},
        }
        PlantillaService._crear_hoja_instrucciones(wb, "Clasificadores de Gasto", campos)
        
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output


plantilla_service = PlantillaService()