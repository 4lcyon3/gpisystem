import io
import pandas as pd
from datetime import datetime
from typing import List, Dict, Any
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
import logging

logger = logging.getLogger(__name__)

# Importar fpdf2 con manejo de errores
try:
    from fpdf import FPDF
    FPDF_AVAILABLE = True
except ImportError:
    logger.error("fpdf2 no está instalado. Ejecuta: pip install fpdf2")
    FPDF_AVAILABLE = False


class ExportService:
    """Servicio centralizado para exportación de datos con diseño profesional"""
    
    COLORS = {
        'primary': '1e3a8a',
        'secondary': '3b82f6',
        'accent': '10b981',
        'warning': 'f59e0b',
        'danger': 'ef4444',
        'light': 'f3f4f6',
        'dark': '1f2937',
    }
    
    @staticmethod
    def to_excel(data: List[Dict], columns: Dict[str, str], sheet_name: str = "Datos") -> io.BytesIO:
        """Exporta datos a Excel con formato profesional"""
        try:
            df = pd.DataFrame(data)
            df = df[[col for col in columns.keys() if col in df.columns]]
            df = df.rename(columns=columns)
            
            output = io.BytesIO()
            
            with pd.ExcelWriter(output, engine='openpyxl', datetime_format='YYYY-MM-DD') as writer:
                df.to_excel(writer, sheet_name=sheet_name, index=False, startrow=1)
                
                workbook = writer.book
                worksheet = writer.sheets[sheet_name]
                
                # Estilos
                header_fill = PatternFill(
                    start_color=ExportService.COLORS['primary'],
                    end_color=ExportService.COLORS['primary'],
                    fill_type='solid'
                )
                header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
                header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                
                thin_border = Border(
                    left=Side(style='thin', color='D1D5DB'),
                    right=Side(style='thin', color='D1D5DB'),
                    top=Side(style='thin', color='D1D5DB'),
                    bottom=Side(style='thin', color='D1D5DB')
                )
                
                alt_fill = PatternFill(
                    start_color=ExportService.COLORS['light'],
                    end_color=ExportService.COLORS['light'],
                    fill_type='solid'
                )
                
                # Aplicar estilos al encabezado (fila 2 porque startrow=1)
                for cell in worksheet[2]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = header_alignment
                    cell.border = thin_border
                
                # Aplicar estilos a los datos
                for row_idx, row in enumerate(worksheet.iter_rows(min_row=3, max_row=worksheet.max_row), start=1):
                    for cell in row:
                        cell.border = thin_border
                        cell.alignment = Alignment(vertical='center', wrap_text=True)
                        
                        # Filas alternas
                        if row_idx % 2 == 0:
                            cell.fill = alt_fill
                        
                        # Formato de números
                        if isinstance(cell.value, (int, float)):
                            cell.number_format = '#,##0.00'
                            cell.alignment = Alignment(horizontal='right', vertical='center')
                
                # Ajustar anchos de columna
                for i, col in enumerate(df.columns, 1):
                    column_letter = get_column_letter(i)
                    max_len = max(
                        df[col].astype(str).apply(len).max() if len(df) > 0 else 0,
                        len(col)
                    ) + 3
                    worksheet.column_dimensions[column_letter].width = min(max_len, 50)
                
                # Congelar paneles
                worksheet.freeze_panes = 'A3'
                
                # Agregar tabla con filtros
                if len(df) > 0:
                    table_ref = f'A2:{get_column_letter(len(df.columns))}{len(df) + 2}'
                    table = Table(displayName=f'Tabla{sheet_name}', ref=table_ref)
                    style = TableStyleInfo(
                        name='TableStyleMedium9',
                        showFirstColumn=False,
                        showLastColumn=False,
                        showRowStripes=True,
                        showColumnStripes=False
                    )
                    table.tableStyleInfo = style
                    worksheet.add_table(table)
                
                # Título del reporte
                title_cell = worksheet['A1']
                title_cell.value = f'Reporte de {sheet_name}'
                title_cell.font = Font(
                    name='Calibri',
                    size=14,
                    bold=True,
                    color=ExportService.COLORS['primary']
                )
                title_cell.alignment = Alignment(horizontal='left', vertical='center')
                
                worksheet.merge_cells(f'A1:{get_column_letter(len(df.columns))}1')
                worksheet.row_dimensions[1].height = 25
            
            output.seek(0)
            return output
            
        except Exception as e:
            logger.error(f"Error generando Excel: {str(e)}")
            raise
    
    @staticmethod
    def to_csv(data: List[Dict], columns: Dict[str, str]) -> io.BytesIO:
        """Exporta datos a CSV con BOM UTF-8"""
        try:
            df = pd.DataFrame(data)
            df = df[[col for col in columns.keys() if col in df.columns]]
            df = df.rename(columns=columns)
            
            output = io.BytesIO()
            output.write(b'\xef\xbb\xbf')  # BOM UTF-8
            output.write(df.to_csv(index=False).encode('utf-8'))
            output.seek(0)
            return output
            
        except Exception as e:
            logger.error(f"Error generando CSV: {str(e)}")
            raise
    
    @staticmethod
    def generate_certificado_pdf(certificado: Dict) -> io.BytesIO:
        """Genera PDF profesional del Certificado de Crédito Presupuestario"""
        if not FPDF_AVAILABLE:
            raise RuntimeError("fpdf2 no está instalado. Ejecuta: pip install fpdf2")
        
        try:
            class CertPDF(FPDF):
                def header(self):
                    self.set_fill_color(30, 58, 138)
                    self.rect(0, 0, 210, 35, 'F')
                    
                    self.set_font('Helvetica', 'B', 18)
                    self.set_text_color(255, 255, 255)
                    self.set_y(8)
                    self.cell(0, 10, 'CERTIFICADO DE CREDITO PRESUPUESTARIO', 0, 1, 'C')
                    
                    self.set_font('Helvetica', '', 11)
                    self.cell(0, 6, f'N° {certificado.get("numero_certificado", "")}', 0, 1, 'C')
                    
                    self.set_draw_color(16, 185, 129)
                    self.set_line_width(0.8)
                    self.line(10, 37, 200, 37)
                    
                    self.ln(12)
                
                def footer(self):
                    self.set_y(-30)
                    self.set_draw_color(200, 200, 200)
                    self.set_line_width(0.3)
                    self.line(10, self.get_y(), 200, self.get_y())
                    self.ln(3)
                    
                    self.set_font('Helvetica', 'I', 8)
                    self.set_text_color(100, 100, 100)
                    self.cell(0, 5, f'Certificado por: {certificado.get("certificador_nombre", "Sistema")}', 0, 1, 'C')
                    self.cell(0, 5, f'Emitido: {datetime.now().strftime("%d/%m/%Y %H:%M:%S")}', 0, 1, 'C')
                    self.cell(0, 5, f'Página {self.page_no()}', 0, 0, 'C')
            
            pdf = CertPDF()
            pdf.add_page()
            pdf.set_auto_page_break(auto=True, margin=25)
            
            # SECCIÓN 1: INFORMACIÓN GENERAL
            pdf.set_font('Helvetica', 'B', 12)
            pdf.set_text_color(30, 58, 138)
            pdf.cell(0, 8, 'INFORMACION GENERAL', 0, 1)
            pdf.set_draw_color(59, 130, 246)
            pdf.set_line_width(0.5)
            pdf.line(10, pdf.get_y(), 100, pdf.get_y())
            pdf.ln(4)
            
            info_data = [
                ('Entidad:', str(certificado.get('entidad_nombre', ''))),
                ('Fecha de Certificación:', str(certificado.get('fecha_certificacion', ''))),
                ('Año Fiscal:', str(certificado.get('anio_fiscal', ''))),
                ('Estado:', str(certificado.get('estado', '')).upper()),
            ]
            
            for label, value in info_data:
                y_before = pdf.get_y()
                pdf.set_font('Helvetica', 'B', 10)
                pdf.cell(50, 7, label, 0, 0)
                pdf.set_font('Helvetica', '', 10)
                pdf.multi_cell(140, 7, value)
                pdf.set_y(max(pdf.get_y(), y_before + 7))
            
            pdf.ln(6)
            
            # SECCIÓN 2: VINCULACIÓN PRESUPUESTAL
            pdf.set_font('Helvetica', 'B', 12)
            pdf.set_text_color(30, 58, 138)
            pdf.cell(0, 8, 'VINCULACION PRESUPUESTAL', 0, 1)
            pdf.set_draw_color(59, 130, 246)
            pdf.line(10, pdf.get_y(), 100, pdf.get_y())
            pdf.ln(4)
            
            pdf.set_font('Helvetica', 'B', 10)
            pdf.cell(45, 7, 'Disponibilidad:', 0, 0)
            pdf.set_font('Helvetica', '', 10)
            pdf.multi_cell(145, 7, str(certificado.get('disponibilidad_numero', '')))
            pdf.ln(6)
            
            # SECCIÓN 3: DETALLE DE MONTOS
            pdf.set_font('Helvetica', 'B', 12)
            pdf.set_text_color(30, 58, 138)
            pdf.cell(0, 8, 'DETALLE DE MONTOS', 0, 1)
            pdf.set_draw_color(59, 130, 246)
            pdf.line(10, pdf.get_y(), 100, pdf.get_y())
            pdf.ln(4)
            
            monto_aprobado = float(certificado.get('disponibilidad_monto_aprobado', 0) or 0)
            monto_certificado = float(certificado.get('monto_certificado', 0) or 0)
            
            # Encabezado de tabla
            pdf.set_fill_color(30, 58, 138)
            pdf.set_text_color(255, 255, 255)
            pdf.set_font('Helvetica', 'B', 11)
            pdf.cell(95, 10, 'Monto Aprobado (Disponibilidad)', 1, 0, 'C', True)
            pdf.cell(95, 10, 'Monto Certificado', 1, 1, 'C', True)
            
            # Datos
            pdf.set_fill_color(240, 248, 255)
            pdf.set_text_color(31, 41, 55)
            pdf.set_font('Helvetica', '', 12)
            pdf.cell(95, 12, f'S/ {monto_aprobado:,.2f}', 1, 0, 'R', True)
            
            pdf.set_fill_color(220, 252, 231)
            pdf.set_font('Helvetica', 'B', 12)
            pdf.cell(95, 12, f'S/ {monto_certificado:,.2f}', 1, 1, 'R', True)
            
            pdf.ln(6)
            
            # SECCIÓN 4: OBSERVACIONES
            if certificado.get('observaciones'):
                pdf.set_font('Helvetica', 'B', 12)
                pdf.set_text_color(30, 58, 138)
                pdf.cell(0, 8, 'OBSERVACIONES', 0, 1)
                pdf.set_draw_color(59, 130, 246)
                pdf.line(10, pdf.get_y(), 100, pdf.get_y())
                pdf.ln(4)
                
                pdf.set_font('Helvetica', '', 10)
                pdf.set_text_color(31, 41, 55)
                pdf.multi_cell(190, 6, str(certificado.get('observaciones', '')))
                pdf.ln(6)
            
            output = io.BytesIO()
            output.write(pdf.output())
            output.seek(0)
            return output
            
        except Exception as e:
            logger.error(f"Error generando PDF de certificado: {str(e)}", exc_info=True)
            raise
    
    @staticmethod
    def generate_modificacion_pdf(modificacion: Dict) -> io.BytesIO:
        """Genera PDF profesional de la Resolución de Modificación Presupuestaria"""
        if not FPDF_AVAILABLE:
            raise RuntimeError("fpdf2 no está instalado. Ejecuta: pip install fpdf2")
        
        try:
            class ModPDF(FPDF):
                def header(self):
                    self.set_fill_color(30, 58, 138)
                    self.rect(0, 0, 210, 35, 'F')
                    
                    self.set_font('Helvetica', 'B', 16)
                    self.set_text_color(255, 255, 255)
                    self.set_y(8)
                    self.cell(0, 10, 'RESOLUCION DE MODIFICACION', 0, 1, 'C')
                    self.cell(0, 8, 'PRESUPUESTARIA', 0, 1, 'C')
                    
                    self.set_font('Helvetica', '', 11)
                    self.cell(0, 6, str(modificacion.get('numero_resolucion', '')), 0, 1, 'C')
                    
                    self.set_draw_color(16, 185, 129)
                    self.set_line_width(0.8)
                    self.line(10, 37, 200, 37)
                    
                    self.ln(12)
                
                def footer(self):
                    self.set_y(-25)
                    self.set_draw_color(200, 200, 200)
                    self.set_line_width(0.3)
                    self.line(10, self.get_y(), 200, self.get_y())
                    self.ln(3)
                    
                    self.set_font('Helvetica', 'I', 8)
                    self.set_text_color(100, 100, 100)
                    self.cell(0, 5, f'Emitido: {datetime.now().strftime("%d/%m/%Y %H:%M:%S")}', 0, 1, 'C')
                    self.cell(0, 5, f'Página {self.page_no()}', 0, 0, 'C')
            
            pdf = ModPDF()
            pdf.add_page()
            pdf.set_auto_page_break(auto=True, margin=25)
            
            # SECCIÓN 1: INFORMACIÓN GENERAL
            pdf.set_font('Helvetica', 'B', 12)
            pdf.set_text_color(30, 58, 138)
            pdf.cell(0, 8, 'INFORMACION GENERAL', 0, 1)
            pdf.set_draw_color(59, 130, 246)
            pdf.set_line_width(0.5)
            pdf.line(10, pdf.get_y(), 100, pdf.get_y())
            pdf.ln(4)
            
            info_data = [
                ('Entidad:', str(modificacion.get('entidad_nombre', ''))),
                ('Fecha de Aprobación:', str(modificacion.get('fecha_aprobacion', ''))),
                ('Tipo de Modificación:', str(modificacion.get('tipo_modificacion', ''))),
                ('Estado:', str(modificacion.get('estado', '')).upper()),
            ]
            
            for label, value in info_data:
                y_before = pdf.get_y()
                pdf.set_font('Helvetica', 'B', 10)
                pdf.cell(50, 7, label, 0, 0)
                pdf.set_font('Helvetica', '', 10)
                pdf.multi_cell(140, 7, value)
                pdf.set_y(max(pdf.get_y(), y_before + 7))
            
            pdf.ln(6)
            
            # SECCIÓN 2: MONTO TOTAL
            pdf.set_font('Helvetica', 'B', 12)
            pdf.set_text_color(30, 58, 138)
            pdf.cell(0, 8, 'MONTO TOTAL DE LA MODIFICACION', 0, 1)
            pdf.set_draw_color(59, 130, 246)
            pdf.line(10, pdf.get_y(), 100, pdf.get_y())
            pdf.ln(4)
            
            pdf.set_fill_color(220, 252, 231)
            pdf.set_draw_color(16, 185, 129)
            pdf.set_line_width(0.8)
            
            monto_total = float(modificacion.get('monto_total', 0) or 0)
            
            pdf.set_font('Helvetica', 'B', 16)
            pdf.set_text_color(31, 41, 55)
            pdf.cell(190, 20, f'S/ {monto_total:,.2f}', 'LTRB', 1, 'C', True)
            
            pdf.ln(6)
            
            # SECCIÓN 3: DESCRIPCIÓN
            if modificacion.get('descripcion'):
                pdf.set_font('Helvetica', 'B', 12)
                pdf.set_text_color(30, 58, 138)
                pdf.cell(0, 8, 'DESCRIPCION Y JUSTIFICACION', 0, 1)
                pdf.set_draw_color(59, 130, 246)
                pdf.line(10, pdf.get_y(), 100, pdf.get_y())
                pdf.ln(4)
                
                pdf.set_font('Helvetica', '', 10)
                pdf.set_text_color(31, 41, 55)
                pdf.set_fill_color(249, 250, 251)
                pdf.multi_cell(190, 6, str(modificacion.get('descripcion', '')), border=1, fill=True)
            
            output = io.BytesIO()
            output.write(pdf.output())
            output.seek(0)
            return output
            
        except Exception as e:
            logger.error(f"Error generando PDF de modificación: {str(e)}", exc_info=True)
            raise


export_service = ExportService()